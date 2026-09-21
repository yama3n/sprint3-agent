import io
import uuid
from pathlib import Path

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.db import AsyncSessionLocal
from app.models.export import Export
from app.models.field import InquiryField, InquiryItemField
from app.models.inquiry import Inquiry, InquiryItem
from app.repositories.field_definition_repository import FieldDefinitionRepository
from app.repositories.inquiry_field_repository import InquiryFieldRepository
from app.repositories.inquiry_item_field_repository import InquiryItemFieldRepository
from app.repositories.inquiry_item_repository import InquiryItemRepository
from app.repositories.inquiry_repository import InquiryRepository


@pytest.fixture
async def db_session() -> AsyncSession:
    async with AsyncSessionLocal() as session:
        yield session
        await session.rollback()


@pytest.fixture
async def created_inquiry_ids():
    ids: list[int] = []
    yield ids
    if not ids:
        return
    async with AsyncSessionLocal() as s:
        item_ids = [
            r[0]
            for r in (
                await s.execute(
                    select(InquiryItem.id).where(InquiryItem.inquiry_id.in_(ids))
                )
            ).fetchall()
        ]
        await s.execute(
            delete(InquiryItemField).where(
                InquiryItemField.inquiry_item_id.in_(item_ids)
            )
        )
        await s.execute(delete(InquiryItem).where(InquiryItem.id.in_(item_ids)))
        await s.execute(delete(InquiryField).where(InquiryField.inquiry_id.in_(ids)))
        await s.execute(delete(Export).where(Export.inquiry_id.in_(ids)))
        await s.execute(delete(Inquiry).where(Inquiry.id.in_(ids)))
        await s.commit()


def auth_headers() -> dict[str, str]:
    return {"Authorization": f"Bearer {settings.MOCK_AUTH_TOKEN}"}


async def _seed(db_session: AsyncSession, created_inquiry_ids: list[int]) -> int:
    inquiry = await InquiryRepository(db_session).add(
        Inquiry(
            inquiry_code=f"INQ-TEST-EXPORT-{uuid.uuid4().hex[:6]}",
            requester="東西石油開発",
            project_name="北海油田 鋼管更新案件",
            status="draft",
        )
    )
    await db_session.flush()
    created_inquiry_ids.append(inquiry.id)

    case_defs = await FieldDefinitionRepository(db_session).list_by_scope("case")
    await InquiryFieldRepository(db_session).add(
        InquiryField(
            inquiry_id=inquiry.id,
            field_definition_id=case_defs[0].id,
            value="東西石油開発株式会社",
            status="ok",
            confirmed_by="user",
        )
    )
    # 記載なし（要確認/missing）→ Excelでは「不明」で出力される
    await InquiryFieldRepository(db_session).add(
        InquiryField(
            inquiry_id=inquiry.id,
            field_definition_id=case_defs[1].id,
            value=None,
            status="review",
            reason_type="missing",
        )
    )

    item = await InquiryItemRepository(db_session).add(
        InquiryItem(inquiry_id=inquiry.id, item_no=1)
    )
    await db_session.flush()
    item_defs = await FieldDefinitionRepository(db_session).list_by_scope("item")
    await InquiryItemFieldRepository(db_session).add(
        InquiryItemField(
            inquiry_item_id=item.id,
            field_definition_id=item_defs[2].id,
            value="L-80",
            status="ok",
            confirmed_by="ai",
        )
    )
    await db_session.commit()
    return inquiry.id


async def test_confirm_requires_auth(client: AsyncClient) -> None:
    resp = await client.post("/api/v1/inquiries/1/confirm")
    assert resp.status_code == 401


async def test_confirm_moves_inquiry_to_final(
    client: AsyncClient, db_session: AsyncSession, created_inquiry_ids: list[int]
) -> None:
    """TEST-19相当: 確定すると status=final になる。"""
    inquiry_id = await _seed(db_session, created_inquiry_ids)

    resp = await client.post(
        f"/api/v1/inquiries/{inquiry_id}/confirm", headers=auth_headers()
    )
    assert resp.status_code == 200
    assert resp.json()["status"] == "final"

    detail = await client.get(f"/api/v1/inquiries/{inquiry_id}", headers=auth_headers())
    assert detail.json()["inquiry"]["status"] == "final"


async def test_confirm_is_idempotent(
    client: AsyncClient, db_session: AsyncSession, created_inquiry_ids: list[int]
) -> None:
    """MVPはロック・確定解除なし: 確定済みへの再確定もエラーにしない（FUNC-09）。"""
    inquiry_id = await _seed(db_session, created_inquiry_ids)
    await client.post(f"/api/v1/inquiries/{inquiry_id}/confirm", headers=auth_headers())
    resp = await client.post(
        f"/api/v1/inquiries/{inquiry_id}/confirm", headers=auth_headers()
    )
    assert resp.status_code == 200


async def test_export_rejects_word_and_pdf_in_scope1(
    client: AsyncClient, db_session: AsyncSession, created_inquiry_ids: list[int]
) -> None:
    """TEST-24相当: Scope 1で有効なのはExcelのみ。"""
    inquiry_id = await _seed(db_session, created_inquiry_ids)

    for fmt in ("word", "pdf"):
        resp = await client.post(
            f"/api/v1/inquiries/{inquiry_id}/exports",
            json={"format": fmt},
            headers=auth_headers(),
        )
        assert resp.status_code == 400
        assert resp.json()["detail"] == "UNSUPPORTED_FORMAT"


async def test_export_draft_uses_draft_suffix(
    client: AsyncClient, db_session: AsyncSession, created_inquiry_ids: list[int]
) -> None:
    inquiry_id = await _seed(db_session, created_inquiry_ids)

    resp = await client.post(
        f"/api/v1/inquiries/{inquiry_id}/exports",
        json={"format": "excel"},
        headers=auth_headers(),
    )
    assert resp.status_code == 200
    assert resp.json()["file_name"].endswith("_draft.xlsx")


async def test_export_after_confirm_uses_final_suffix_and_downloads(
    client: AsyncClient, db_session: AsyncSession, created_inquiry_ids: list[int]
) -> None:
    """TEST-24相当: 確定後の出力は _final 付き。ダウンロード内容が確定済みJSONと一致する。"""
    inquiry_id = await _seed(db_session, created_inquiry_ids)
    await client.post(f"/api/v1/inquiries/{inquiry_id}/confirm", headers=auth_headers())

    resp = await client.post(
        f"/api/v1/inquiries/{inquiry_id}/exports",
        json={"format": "excel"},
        headers=auth_headers(),
    )
    body = resp.json()
    assert body["file_name"].endswith("_final.xlsx")

    download = await client.get(
        f"/api/v1/inquiries/{inquiry_id}/exports/{body['export_id']}/download",
        headers=auth_headers(),
    )
    assert download.status_code == 200
    assert "attachment" in download.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(download.content))
    texts = [
        str(cell.value)
        for row in workbook["Item List"].iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert "東西石油開発株式会社" in texts  # 確定済みの値
    assert "L-80" in texts  # 品目一覧
    assert "不明" in texts  # missing項目は「不明」（FUNC-05）
    assert "要確認項目" in texts  # FUNC-05の表示順の最後のセクション


async def test_download_404_for_unknown_export(
    client: AsyncClient, db_session: AsyncSession, created_inquiry_ids: list[int]
) -> None:
    inquiry_id = await _seed(db_session, created_inquiry_ids)
    resp = await client.get(
        f"/api/v1/inquiries/{inquiry_id}/exports/999999/download",
        headers=auth_headers(),
    )
    assert resp.status_code == 404
    assert resp.json()["detail"] == "EXPORT_NOT_FOUND"


async def test_export_missing_items_do_not_block_output(
    client: AsyncClient, db_session: AsyncSession, created_inquiry_ids: list[int]
) -> None:
    """FUNC-05受入基準: 要確認/missingが残っていても出力自体は可能。"""
    inquiry_id = await _seed(db_session, created_inquiry_ids)
    resp = await client.post(
        f"/api/v1/inquiries/{inquiry_id}/exports",
        json={"format": "excel"},
        headers=auth_headers(),
    )
    assert resp.status_code == 200
    assert Path(settings.STORAGE_DIR, "exports", str(inquiry_id)).exists()
