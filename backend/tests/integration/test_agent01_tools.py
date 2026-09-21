import json
import uuid

import openpyxl
import pytest
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.agent.agent01 import state
from app.agent.agent01.tools import (
    emit_extraction_result,
    extract_case_fields,
    extract_item_fields,
    extract_supplementary_notes,
    parse_excel,
)
from app.core.db import AsyncSessionLocal
from app.models.agent_run import AgentRun, ExtractionResult, ParsedDocument
from app.models.field import InquiryField
from app.models.inquiry import Inquiry, InquiryFile
from app.repositories.agent_run_repository import AgentRunRepository
from app.repositories.inquiry_file_repository import InquiryFileRepository
from app.repositories.inquiry_repository import InquiryRepository


@pytest.fixture
async def db_session() -> AsyncSession:
    async with AsyncSessionLocal() as session:
        yield session
        await session.rollback()


@pytest.fixture
async def created_inquiry_ids():
    ids: list[int] = []
    yield ids
    if not ids:
        return
    async with AsyncSessionLocal() as cleanup_session:
        file_ids_result = await cleanup_session.execute(
            InquiryFile.__table__.select().where(InquiryFile.inquiry_id.in_(ids))
        )
        file_ids = [row.id for row in file_ids_result.fetchall()]
        await cleanup_session.execute(
            delete(ParsedDocument).where(ParsedDocument.file_id.in_(file_ids))
        )
        await cleanup_session.execute(
            delete(ExtractionResult).where(ExtractionResult.inquiry_id.in_(ids))
        )
        await cleanup_session.execute(
            delete(AgentRun).where(AgentRun.inquiry_id.in_(ids))
        )
        await cleanup_session.execute(
            delete(InquiryField).where(InquiryField.inquiry_id.in_(ids))
        )
        await cleanup_session.execute(
            delete(InquiryFile).where(InquiryFile.inquiry_id.in_(ids))
        )
        await cleanup_session.execute(delete(Inquiry).where(Inquiry.id.in_(ids)))
        await cleanup_session.commit()
    for inquiry_id in ids:
        state.clear_state(inquiry_id)


async def _make_inquiry_with_excel_file(
    db_session: AsyncSession, tmp_path, created_inquiry_ids: list[int]
) -> tuple[int, int]:
    inquiry = await InquiryRepository(db_session).add(
        Inquiry(inquiry_code=f"INQ-TEST-AGENT01-{uuid.uuid4().hex[:8]}", status="draft")
    )
    await db_session.flush()
    created_inquiry_ids.append(inquiry.id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "発注元: 東西石油開発株式会社"
    ws["A2"] = "案件名: 北海油田 鋼管更新案件"
    path = tmp_path / "order.xlsx"
    wb.save(path)

    inquiry_file = await InquiryFileRepository(db_session).add(
        InquiryFile(
            inquiry_id=inquiry.id,
            file_name="order.xlsx",
            file_type="xlsx",
            storage_path=str(path),
        )
    )
    await db_session.commit()
    return inquiry.id, inquiry_file.id


async def test_parse_excel_tool_persists_parsed_document_and_returns_content(
    db_session: AsyncSession, tmp_path, created_inquiry_ids: list[int]
) -> None:
    _, file_id = await _make_inquiry_with_excel_file(
        db_session, tmp_path, created_inquiry_ids
    )

    response = await parse_excel.handler({"file_id": file_id})

    assert response.get("is_error") is not True
    payload = json.loads(response["content"][0]["text"])
    assert payload["file_name"] == "order.xlsx"
    assert any("東西石油開発" in row["text"] for row in payload["sheets"])


async def test_parse_excel_tool_returns_error_for_unknown_file() -> None:
    response = await parse_excel.handler({"file_id": 999999999})
    assert response.get("is_error") is True


async def test_extraction_pipeline_end_to_end_persists_extraction_result(
    db_session: AsyncSession, tmp_path, created_inquiry_ids: list[int]
) -> None:
    inquiry_id, _ = await _make_inquiry_with_excel_file(
        db_session, tmp_path, created_inquiry_ids
    )
    agent_run = await AgentRunRepository(db_session).add(
        AgentRun(
            inquiry_id=inquiry_id,
            agent_name="agent01_extraction",
            trigger="new_upload",
            status="running",
            stage="extracting",
        )
    )
    await db_session.commit()

    case_result = await extract_case_fields.handler(
        {
            "inquiry_id": inquiry_id,
            "candidates": [
                {
                    "field_id": "requester",
                    "value": "東西石油開発株式会社",
                    "source_type": "excel",
                    "source_file": "order.xlsx",
                    "source_location": "A1",
                    "quoted_text": "発注元: 東西石油開発株式会社",
                }
            ],
        }
    )
    assert case_result.get("is_error") is not True

    item_result = await extract_item_fields.handler(
        {"inquiry_id": inquiry_id, "items": [{"item_no": 1, "candidates": []}]}
    )
    assert item_result.get("is_error") is not True

    notes_result = await extract_supplementary_notes.handler(
        {"inquiry_id": inquiry_id, "case_notes": [], "item_notes": []}
    )
    assert notes_result.get("is_error") is not True

    emit_result = await emit_extraction_result.handler(
        {
            "inquiry_id": inquiry_id,
            "agent_run_id": agent_run.id,
            "source_files": ["order.xlsx"],
        }
    )
    assert emit_result.get("is_error") is not True
    payload = json.loads(emit_result["content"][0]["text"])
    assert payload["saved"] is True

    async with AsyncSessionLocal() as verify_session:
        from sqlalchemy import select

        result = await verify_session.execute(
            select(ExtractionResult).where(ExtractionResult.inquiry_id == inquiry_id)
        )
        saved = result.scalar_one()
        assert (
            saved.payload["case_fields"]["requester"][0]["value"]
            == "東西石油開発株式会社"
        )
        assert "1" in saved.payload["item_fields"]  # JSON永続化でitem_noは文字列キー化

        run = await verify_session.get(AgentRun, agent_run.id)
        assert run.status == "succeeded"
        assert run.stage == "completed"
        assert run.progress_percent == 55


async def test_emit_extraction_result_fails_completion_check_when_case_fields_not_extracted(
    db_session: AsyncSession, tmp_path, created_inquiry_ids: list[int]
) -> None:
    inquiry_id, _ = await _make_inquiry_with_excel_file(
        db_session, tmp_path, created_inquiry_ids
    )
    agent_run = await AgentRunRepository(db_session).add(
        AgentRun(
            inquiry_id=inquiry_id,
            agent_name="agent01_extraction",
            trigger="new_upload",
            status="running",
            stage="extracting",
        )
    )
    await db_session.commit()

    response = await emit_extraction_result.handler(
        {
            "inquiry_id": inquiry_id,
            "agent_run_id": agent_run.id,
            "source_files": ["order.xlsx"],
        }
    )

    assert response.get("is_error") is True
