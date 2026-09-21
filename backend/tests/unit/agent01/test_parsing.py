import email.message
from pathlib import Path

import openpyxl
from fpdf import FPDF

from app.agent.agent01.parsing import parse_eml_file, parse_excel_file, parse_pdf_file


def test_parse_excel_file_extracts_cells_with_coordinates(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "発注リスト"
    ws["A1"] = "品名"
    ws["B1"] = "数量"
    ws["A2"] = "鋼管"
    ws["B2"] = 100
    path = tmp_path / "order.xlsx"
    wb.save(path)

    rows = parse_excel_file(path)

    assert {"sheet": "発注リスト", "cell": "A1", "text": "品名"} in rows
    assert {"sheet": "発注リスト", "cell": "B2", "text": "100"} in rows
    # 空セルは含めない
    assert not any(r["cell"] == "C1" for r in rows)


def test_parse_excel_file_handles_multiple_sheets(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "one"
    wb.create_sheet("Sheet2")["A1"] = "two"
    path = tmp_path / "multi.xlsx"
    wb.save(path)

    rows = parse_excel_file(path)

    sheets = {r["sheet"] for r in rows}
    assert sheets == {"Sheet1", "Sheet2"}


def test_parse_pdf_file_extracts_text_per_page(tmp_path: Path) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Quantity: 320 pieces")
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Delivery: Norway")
    path = tmp_path / "quote.pdf"
    pdf.output(str(path))

    pages = parse_pdf_file(path)

    assert len(pages) == 2
    assert pages[0]["page"] == 1
    assert "320" in pages[0]["text"]
    assert pages[1]["page"] == 2
    assert "Norway" in pages[1]["text"]


def test_parse_eml_file_splits_subject_and_body_paragraphs(tmp_path: Path) -> None:
    msg = email.message.EmailMessage()
    msg["Subject"] = "北海油田案件の引合"
    msg["From"] = "customer@example.com"
    msg["To"] = "sales@toseki-steel.co.jp"
    msg.set_content(
        "お世話になっております。\n\n"
        "数量は240本でお願いします。\n\n"
        "先ほどの数量240本ではなく320本でお願いします。"
    )
    path = tmp_path / "inquiry.eml"
    path.write_bytes(bytes(msg))

    sections = parse_eml_file(path)

    assert sections[0] == {"section": "subject", "text": "北海油田案件の引合"}
    body_sections = [s for s in sections if s["section"] != "subject"]
    assert len(body_sections) == 3
    assert "320本" in body_sections[2]["text"]
