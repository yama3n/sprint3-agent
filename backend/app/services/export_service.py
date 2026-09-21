"""FUNC-09 成果物確定 / FUNC-05 成果物生成（Excel Item List）。

エージェントではなく決定的な生成処理（agent-plan.md 5章「非エージェント処理」）。
構造化JSONから一方向に生成し、生成物を読み戻して反映することはしない。
"""

import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.v1.schemas.inquiry import (
    ConfirmResponse,
    ExportResponse,
    InquiryDetailResponse,
)
from app.core.config import settings
from app.models.export import Export
from app.repositories.export_repository import ExportRepository
from app.repositories.inquiry_repository import InquiryRepository
from app.services.inquiry_service import InquiryNotFoundError, get_inquiry_detail

# FUNC-05: missing（記載なし）はExcel上「不明」と表示する（画面のステータス表現とは別概念）
MISSING_CELL_TEXT = "不明"


class UnsupportedExportFormatError(Exception):
    pass


class ExportNotFoundError(Exception):
    pass


async def confirm_inquiry(session: AsyncSession, inquiry_id: int) -> ConfirmResponse:
    """FUNC-09: 構造化JSONを確定済み（final）にする。

    MVPではロック・確定解除・版管理は行わない（Scope 2）。確定済みへの再確定もエラーにしない。
    """
    inquiry = await InquiryRepository(session).get(inquiry_id)
    if inquiry is None:
        raise InquiryNotFoundError(inquiry_id)

    inquiry.status = "final"
    await session.commit()
    await session.refresh(inquiry)

    return ConfirmResponse(
        inquiry_id=inquiry_id, status="final", confirmed_at=inquiry.updated_at
    )


def _cell_text(value: str | None, status: str) -> str:
    """要確認かつ値なし（missing相当）は「不明」を出力する（FUNC-05受入基準）。"""
    if value:
        return value
    return MISSING_CELL_TEXT if status == "review" else ""


def build_workbook(detail: InquiryDetailResponse) -> Workbook:
    """構造化JSONからItem List（Excel）を生成する。

    FUNC-05の表示順（依頼元企業 → 案件情報 → 品目一覧 → その他特記事項 → 要確認項目）に従う。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Item List"
    bold = Font(bold=True)
    row = 1

    def write_heading(text: str) -> None:
        nonlocal row
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font = bold
        row += 1

    # 1. 依頼元企業 / 2. 案件情報（A項目14件をdisplay_order順に）
    write_heading("引合情報")
    sheet.cell(row=row, column=1, value="引合ID").font = bold
    sheet.cell(row=row, column=2, value=detail.inquiry.inquiry_code)
    row += 1
    for field in detail.case_fields:
        sheet.cell(row=row, column=1, value=field.label).font = bold
        sheet.cell(row=row, column=2, value=_cell_text(field.value, field.status))
        if field.is_web_supplemented:
            sheet.cell(row=row, column=3, value="Web補完")
        row += 1
    row += 1

    # 3. 品目一覧（B項目8件 × 品目数）
    write_heading("品目一覧")
    if detail.items:
        headers = ["#"] + [f.label for f in detail.items[0].fields]
        for column, header in enumerate(headers, start=1):
            sheet.cell(row=row, column=column, value=header).font = bold
        row += 1
        for item in detail.items:
            sheet.cell(row=row, column=1, value=item.item_no)
            for column, field in enumerate(item.fields, start=2):
                sheet.cell(
                    row=row, column=column, value=_cell_text(field.value, field.status)
                )
            row += 1
    row += 1

    # 4. その他特記事項（案件全体 / 品目固有）
    write_heading("その他特記事項")
    for note in detail.case_notes:
        sheet.cell(row=row, column=1, value=note.content)
        sheet.cell(row=row, column=2, value=note.source_file or "")
        row += 1
    for item in detail.items:
        for note in item.notes:
            sheet.cell(row=row, column=1, value=f"品目{item.item_no}: {note.content}")
            sheet.cell(row=row, column=2, value=note.source_file or "")
            row += 1
    row += 1

    # 5. 要確認項目（担当者が確認すべき箇所の一覧）
    write_heading("要確認項目")
    review_rows = [(f.label, None) for f in detail.case_fields if f.status == "review"]
    review_rows += [
        (f.label, item.item_no)
        for item in detail.items
        for f in item.fields
        if f.status == "review"
    ]
    for label, item_no in review_rows:
        sheet.cell(
            row=row,
            column=1,
            value=label if item_no is None else f"品目{item_no}: {label}",
        )
        row += 1

    return workbook


def _export_file_name(inquiry_code: str, state: str) -> str:
    """FUNC-05/FUNC-09: 確定前は _draft、確定済みは _final を付与する。"""
    suffix = "final" if state == "final" else "draft"
    return f"{inquiry_code}_{suffix}.xlsx"


async def create_export(
    session: AsyncSession, inquiry_id: int, export_format: str
) -> ExportResponse:
    """FUNC-05: 構造化JSONからExcelを生成し、exports に記録する。"""
    if export_format != "excel":
        # Scope 1 で有効なのは Excel のみ（Word/PDF は Scope 2）
        raise UnsupportedExportFormatError(export_format)

    inquiry = await InquiryRepository(session).get(inquiry_id)
    if inquiry is None:
        raise InquiryNotFoundError(inquiry_id)

    detail = await get_inquiry_detail(session, inquiry_id)
    workbook = build_workbook(detail)

    export_dir = Path(settings.STORAGE_DIR) / "exports" / str(inquiry_id)
    export_dir.mkdir(parents=True, exist_ok=True)
    file_name = _export_file_name(inquiry.inquiry_code, inquiry.status)
    timestamp = datetime.datetime.now(datetime.UTC).strftime("%Y%m%d%H%M%S")
    storage_path = export_dir / f"{timestamp}_{file_name}"
    workbook.save(storage_path)

    export = await ExportRepository(session).add(
        Export(
            inquiry_id=inquiry_id,
            format="excel",
            state=inquiry.status,
            file_name=file_name,
            storage_path=str(storage_path),
        )
    )
    await session.commit()

    return ExportResponse(export_id=export.id, file_name=file_name)


async def get_export_file(
    session: AsyncSession, inquiry_id: int, export_id: int
) -> tuple[Path, str]:
    """GET /inquiries/{id}/exports/{export_id}/download 用にファイルパスと表示名を返す。"""
    inquiry = await InquiryRepository(session).get(inquiry_id)
    if inquiry is None:
        raise InquiryNotFoundError(inquiry_id)

    export = await ExportRepository(session).get(export_id)
    if export is None or export.inquiry_id != inquiry_id:
        raise ExportNotFoundError(export_id)

    path = Path(export.storage_path)
    if not path.exists():
        raise ExportNotFoundError(export_id)
    return path, export.file_name
